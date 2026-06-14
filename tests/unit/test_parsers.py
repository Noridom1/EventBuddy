from io import BytesIO

from eventbuddy.ingestion.parsers import parse


def _xlsx_bytes(rows):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes(paragraphs):
    import docx

    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_parse_xlsx_headers_to_rows():
    content = _xlsx_bytes([["Email", "Rating", "Comment"],
                           ["a@x.com", 5, "great"],
                           ["b@x.com", 2, "long"]])
    doc = parse("responses.xlsx", content)
    assert doc.kind == "xlsx"
    assert len(doc.rows) == 2
    assert doc.rows[0] == {"Email": "a@x.com", "Rating": 5, "Comment": "great"}
    assert "a@x.com" in doc.text


def test_parse_xlsx_skips_blank_rows():
    content = _xlsx_bytes([["A", "B"], [None, None], ["x", "y"]])
    doc = parse("f.xlsx", content)
    assert len(doc.rows) == 1


def test_parse_docx_extracts_text():
    doc = parse("plan.docx", _docx_bytes(["Hello world", "Second line"]))
    assert doc.kind == "docx"
    assert "Hello world" in doc.text
    assert "Second line" in doc.text


def test_parse_unsupported_extension():
    assert parse("notes.rtf", b"junk").kind == "unsupported"


def test_parse_pdf_garbage_degrades_to_unsupported():
    # Not a real PDF — the guarded parser must degrade, not raise.
    assert parse("broken.pdf", b"not a pdf").kind == "unsupported"


# --- Impl 5: image / image-PDF -------------------------------------------------------------

def test_parse_image_returns_image_kind_with_bytes():
    doc = parse("flyer.png", b"\x89PNG-bytes")
    assert doc.kind == "image"
    assert doc.raw_bytes == b"\x89PNG-bytes"
    assert doc.mime == "image/png"


def test_parse_jpeg_maps_mime():
    assert parse("photo.JPG", b"jpegbytes").mime == "image/jpeg"


def test_parse_plaintext():
    doc = parse("notes.txt", b"Hello\nWorld")
    assert doc.kind == "txt"
    assert doc.text == "Hello\nWorld"


def test_empty_text_pdf_flagged_image_pdf():
    import pypdf

    buf = BytesIO()
    w = pypdf.PdfWriter()
    w.add_blank_page(width=200, height=200)  # a page with no extractable text
    w.write(buf)
    doc = parse("scan.pdf", buf.getvalue())
    assert doc.kind == "image_pdf"
    assert doc.raw_bytes is not None
    assert doc.mime == "application/pdf"


# --- Impl 4: CSV / TSV ---------------------------------------------------------------------

def test_parse_csv_headers_to_rows():
    content = b"Name,Email,Registered\nAlice,a@x.com,Yes\nBob,b@x.com,No\n"
    doc = parse("roster.csv", content)
    assert doc.kind == "csv"
    assert len(doc.rows) == 2
    assert doc.rows[0] == {"Name": "Alice", "Email": "a@x.com", "Registered": "Yes"}
    assert "b@x.com" in doc.text


def test_parse_csv_handles_bom_and_semicolon_delimiter():
    content = "﻿Name;Email\nAlice;a@x.com\n".encode()
    doc = parse("roster.csv", content)
    assert doc.kind == "csv"
    assert doc.rows[0]["Email"] == "a@x.com"  # BOM stripped, ';' sniffed


def test_parse_tsv_forces_tab_delimiter():
    content = b"Name\tEmail\nAlice\ta@x.com\n"
    doc = parse("roster.tsv", content)
    assert doc.kind == "csv"
    assert doc.rows[0] == {"Name": "Alice", "Email": "a@x.com"}


def test_parse_csv_skips_blank_rows():
    content = b"A,B\n,\nx,y\n"
    doc = parse("f.csv", content)
    assert len(doc.rows) == 1
